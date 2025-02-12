from flask import Flask, render_template, request, send_file
import pandas as pd
import os
from datetime import datetime
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import nsmap
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

app = Flask(__name__)
UPLOAD_FOLDER = 'static/uploads'
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
EXCEL_FILE = 'data.xlsx'
LOGO_PATH = 'static/easybuylogo.png'

@app.route('/')
def index():
    return render_template('form.html')

@app.route('/generate_report/<po_number>')
def generate_report(po_number):
    try:
        df = pd.read_excel(EXCEL_FILE)
        row = df[df['PO Number'] == int(po_number)].to_dict(orient='records')
        if not row:
            return "PO Number not found."
        row = row[0]
    except Exception as e:
        return f"Error: {e}"

    doc = Document()
    
    # Add Logo
    if os.path.exists(LOGO_PATH):
        doc.add_picture(LOGO_PATH, width=Inches(2))
        last_paragraph = doc.paragraphs[-1]
        last_paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    
    doc.add_heading('Fabric Inspection Report', level=1)
    
    # Create Table for Key Information
    table = doc.add_table(rows=7, cols=2)
    table.style = 'Table Grid'
    data_pairs = [
        ('Date', row['Date']),
        ('Vendor', row['Vendor']),
        ('PO Number', row['PO Number']),
        ('Order Quantity', row['Order Qty']),
        ('Offer Quantity', row['Offer Qty']),
        ('Style', row['Style']),
        ('Colour', row['Colour'])
    ]
    
    for i, (key, value) in enumerate(data_pairs):
        table.cell(i, 0).text = key
        table.cell(i, 1).text = str(value)
    
    doc.add_paragraph('\n')
    
    # Defect Details in Rows
    doc.add_heading('Defect Details', level=2)
    defect_table = doc.add_table(rows=6, cols=2)
    defect_table.style = 'Table Grid'
    defect_data = [
        ('Type of Major Defect', row['Type of Major Defect']),
        ('Count of Major Defect', row['Count of Major Defect']),
        ('Type of Minor Defect', row['Type of Minor Defect']),
        ('Count of Minor Defect', row['Count of Minor Defect']),
        ('Inspection Result', row['Inspection Result']),
        ('QA Remarks', row['QA Remarks'])
    ]
    
    for i, (label, value) in enumerate(defect_data):
        defect_table.cell(i, 0).text = label
        defect_table.cell(i, 1).text = str(value)
    
    doc.add_page_break()  # Move images to a new page
    
    # Images Section
    doc.add_heading('Inspection Images', level=2)
    image_fields = [
        ('Bulk vs PP', 'bulk_vs_pp'),
        ('Barcode Tag & Washcare', 'barcode_washcare'),
        ('Bulk vs HLP/GPT Report', 'bulk_vs_hlp'),
        ('Measurement Chart', 'measurement_chart'),
        ('Bulk Carton/Carton Opening/Carton Marking', 'bulk_carton'),
        ('Defect Picture', 'defect_picture')
    ]
    
    image_table = doc.add_table(rows=3, cols=2)
    image_table.style = 'Table Grid'
    
    for i, (label, field) in enumerate(image_fields):
        if row[field]:
            cell = image_table.cell(i // 2, i % 2)
            cell.text = label
            paragraph = cell.add_paragraph()
            run = paragraph.add_run()
            run.add_picture(row[field], width=Inches(2))
    
    # Save and Allow Download
    report_path = f"static/reports/Report_{po_number}.docx"
    os.makedirs("static/reports", exist_ok=True)
    doc.save(report_path)
    
    return send_file(report_path, as_attachment=True)

if __name__ == '__main__':
    app.run(debug=True)